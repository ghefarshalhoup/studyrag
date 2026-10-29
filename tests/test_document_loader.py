import io
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from langchain_core.documents import Document

from src.document_loader import chunk_documents, load_document

DOCUMENTS = Path(__file__).resolve().parents[1] / "evaluation" / "documents"


def test_pdf_pages_keep_source_and_page_number():
    with open(DOCUMENTS / "ground_truth.pdf", "rb") as f:
        pages = load_document(f, "ground_truth.pdf")

    assert pages
    assert pages[0].metadata["source"] == "ground_truth.pdf"
    assert pages[0].metadata["page"] == 1
    assert "Hypothesis testing" in " ".join(p.page_content for p in pages)


def test_docx_is_split_into_estimated_pages():
    docx = DocxDocument()
    for _ in range(8):
        docx.add_paragraph("x" * 1000)  # about 3000 characters per estimated page
    buffer = io.BytesIO()
    docx.save(buffer)
    buffer.seek(0)

    pages = load_document(buffer, "notes.docx")

    assert [p.metadata["page"] for p in pages] == [1, 2, 3]
    assert all(p.metadata["total_pages"] == 3 for p in pages)


def test_unsupported_file_type_is_rejected():
    with pytest.raises(ValueError, match="Unsupported file type"):
        load_document(io.BytesIO(b""), "slides.pptx")


def test_chunks_inherit_metadata_and_are_numbered():
    page = Document(page_content="word " * 600, metadata={"source": "a.pdf", "page": 7})

    chunks = chunk_documents([page])

    assert len(chunks) > 1
    assert all(c.metadata["source"] == "a.pdf" and c.metadata["page"] == 7 for c in chunks)
    assert [c.metadata["chunk_index"] for c in chunks] == list(range(len(chunks)))
    assert all(len(c.page_content) <= 1000 for c in chunks)
