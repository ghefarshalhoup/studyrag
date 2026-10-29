"""Document loading and chunking for PDF and DOCX files."""

import tempfile
from pathlib import Path
from typing import BinaryIO

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from pypdf import PdfReader
from docx import Document as DocxDocument

from src.config import settings


def load_pdf(file: BinaryIO, filename: str) -> list[Document]:
    """Load a PDF file and return documents with page metadata."""
    documents = []

    # Save uploaded file to temp location for processing
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp.write(file.read())
        tmp_path = tmp.name

    try:
        reader = PdfReader(tmp_path)

        for page_num, page in enumerate(reader.pages, start=1):
            text = page.extract_text()
            if text and text.strip():
                documents.append(
                    Document(
                        page_content=text,
                        metadata={
                            "source": filename,
                            "page": page_num,
                            "total_pages": len(reader.pages),
                            "file_type": "pdf",
                        },
                    )
                )
    finally:
        Path(tmp_path).unlink(missing_ok=True)

    return documents


def load_docx(file: BinaryIO, filename: str) -> list[Document]:
    """Load a DOCX file and return documents with page metadata."""
    documents = []

    # Save uploaded file to temp location for processing
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp.write(file.read())
        tmp_path = tmp.name

    try:
        doc = DocxDocument(tmp_path)

        # DOCX doesn't have native page numbers, so we estimate based on paragraphs
        # Group paragraphs into "pages" of approximately 3000 characters
        current_page = 1
        current_content = []
        current_length = 0
        chars_per_page = 3000

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            current_content.append(text)
            current_length += len(text)

            if current_length >= chars_per_page:
                documents.append(
                    Document(
                        page_content="\n".join(current_content),
                        metadata={
                            "source": filename,
                            "page": current_page,
                            "file_type": "docx",
                        },
                    )
                )
                current_page += 1
                current_content = []
                current_length = 0

        # Add remaining content
        if current_content:
            documents.append(
                Document(
                    page_content="\n".join(current_content),
                    metadata={
                        "source": filename,
                        "page": current_page,
                        "file_type": "docx",
                    },
                )
            )

        # Update total pages in metadata
        total_pages = current_page
        for doc in documents:
            doc.metadata["total_pages"] = total_pages

    finally:
        Path(tmp_path).unlink(missing_ok=True)

    return documents


def load_document(file: BinaryIO, filename: str) -> list[Document]:
    """Load a document based on file extension."""
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        return load_pdf(file, filename)
    elif extension == ".docx":
        return load_docx(file, filename)
    else:
        raise ValueError(f"Unsupported file type: {extension}")


def chunk_documents(documents: list[Document]) -> list[Document]:
    """Split documents into smaller chunks while preserving metadata."""
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        length_function=len,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    chunked_docs = []

    for doc in documents:
        chunks = text_splitter.split_text(doc.page_content)

        for i, chunk in enumerate(chunks):
            chunked_docs.append(
                Document(
                    page_content=chunk,
                    metadata={
                        **doc.metadata,
                        "chunk_index": i,
                        "total_chunks": len(chunks),
                    },
                )
            )

    return chunked_docs


def process_uploaded_files(files: list[tuple[BinaryIO, str]]) -> list[Document]:
    """Process multiple uploaded files and return chunked documents."""
    all_documents = []

    for file, filename in files:
        # Reset file pointer
        file.seek(0)

        # Load document
        documents = load_document(file, filename)

        # Chunk documents
        chunked = chunk_documents(documents)

        all_documents.extend(chunked)

    return all_documents
